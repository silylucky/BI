from __future__ import annotations

import os
import time
from io import BytesIO
from pathlib import Path
from typing import Any

import httpx

from app.datasources.dialects.base import ColumnInfo, SchemaInfo, TableInfo, TestConnectionResult
from app.datasources.dialects.errors import (
    FILE_DRIVER_MISSING,
    FILE_EXTENSION_DENIED,
    FILE_NOT_FOUND,
    FILE_PATH_TRAVERSAL,
    map_file_error,
)
from app.query.native.guard import guard_native_injection

ALLOWED_EXCEL_SUFFIX = ".xlsx"


def _resolve_file_path(host: str) -> tuple[Path | None, str | None, bool]:
    if host.startswith("https://"):
        return None, None, True
    if ".." in host.replace("\\", "/").split("/"):
        return None, FILE_PATH_TRAVERSAL, False
    path = Path(host).expanduser().resolve()
    if path.suffix.lower() != ALLOWED_EXCEL_SUFFIX:
        return None, FILE_EXTENSION_DENIED, False
    if not path.is_file() or not os.access(path, os.R_OK):
        return None, FILE_NOT_FOUND, False
    return path, None, False


def _open_workbook(connection: Any) -> Any:
    if isinstance(connection, dict) and connection.get("remote"):
        url = str(connection["remote"])
        with httpx.Client(timeout=30.0, follow_redirects=True) as client:
            resp = client.get(url)
            resp.raise_for_status()
        from openpyxl import load_workbook

        return load_workbook(BytesIO(resp.content), read_only=True, data_only=True)
    return connection


class ExcelConnector:
    type = "excel"
    category = "file"
    capabilities = ("connectivity_test", "schema_browser", "native_query")
    display_name = "Excel"

    def test_connection(self, **kwargs: Any) -> TestConnectionResult:
        started = time.perf_counter()
        path, err, is_remote = _resolve_file_path(kwargs["host"])
        if is_remote:
            try:
                with httpx.Client(timeout=float(kwargs.get("timeout_sec", 5.0)), follow_redirects=True) as client:
                    resp = client.head(kwargs["host"])
                if not resp.is_success:
                    code, detail = map_file_error(Exception(f"remote http {resp.status_code}"))
                    return TestConnectionResult(ok=False, message=f"[{code}] {detail}", latency_ms=0, code=code)
            except Exception as exc:
                code, detail = map_file_error(exc)
                return TestConnectionResult(ok=False, message=f"[{code}] {detail}", latency_ms=0, code=code)
            return TestConnectionResult(
                ok=True,
                message="Remote file reachable",
                latency_ms=int((time.perf_counter() - started) * 1000),
                code=None,
            )
        if err:
            return TestConnectionResult(ok=False, message=f"[{err}] invalid path", latency_ms=0, code=err)
        try:
            from openpyxl import load_workbook

            wb = load_workbook(path, read_only=True, data_only=True)
            wb.close()
        except ImportError:
            return TestConnectionResult(
                ok=False,
                message=f"[{FILE_DRIVER_MISSING}] openpyxl not installed",
                latency_ms=0,
                code=FILE_DRIVER_MISSING,
            )
        except Exception as exc:
            code, detail = map_file_error(exc)
            return TestConnectionResult(ok=False, message=f"[{code}] {detail}", latency_ms=0, code=code)
        return TestConnectionResult(
            ok=True,
            message="Connection successful",
            latency_ms=int((time.perf_counter() - started) * 1000),
            code=None,
        )

    def open_connection(self, **kwargs: Any) -> Any:
        path, err, is_remote = _resolve_file_path(kwargs["host"])
        if is_remote:
            return {"remote": kwargs["host"]}
        if err or path is None:
            raise ValueError(err or FILE_NOT_FOUND)
        from openpyxl import load_workbook

        return load_workbook(path, read_only=True, data_only=True)

    def list_schemas(self, connection: Any) -> list[SchemaInfo]:
        return [SchemaInfo(name="workbook")]

    def list_tables(self, connection: Any, schema: str) -> list[TableInfo]:
        if hasattr(connection, "sheetnames"):
            return [TableInfo(name=n, type="SHEET") for n in connection.sheetnames]
        return [TableInfo(name="Sheet1", type="SHEET")]

    def list_columns(self, connection: Any, schema: str, table: str) -> list[ColumnInfo]:
        workbook = _open_workbook(connection)
        sheet_name = table or (workbook.sheetnames[0] if hasattr(workbook, "sheetnames") else "Sheet1")
        ws = workbook[sheet_name]
        row = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), ())
        return [
            ColumnInfo(name=str(c) if c is not None else f"col_{idx}", data_type="string", nullable=True)
            for idx, c in enumerate(row)
        ]

    def execute_native_query(
        self,
        connection: Any,
        *,
        body: dict,
        limit: int,
        offset: int = 0,
        database: str | None = None,
    ) -> tuple[list[str], list[list], bool]:
        guard_native_injection(body)
        workbook = _open_workbook(connection)
        sheet = body.get("table") or database or "Sheet1"
        ws = workbook[str(sheet)]
        rows_iter = ws.iter_rows(values_only=True)
        header_row = next(rows_iter, ())
        header = [str(c) if c is not None else f"col_{idx}" for idx, c in enumerate(header_row)]
        data = [list(r) for r in rows_iter if any(r)]
        data = data[offset : offset + limit + 1]
        truncated = len(data) > limit
        return header, data[:limit], truncated
